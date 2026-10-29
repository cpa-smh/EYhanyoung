from copy import deepcopy
from decimal import Decimal as D
import json
from zipfile import ZipFile
import pytest
from openpyxl import load_workbook
from PIL import Image
from pipeline import analyze, load_snapshot, save_snapshot
from main import main


def test_pipeline_observation_counts_and_three_year_results(analysis):
    assert len(analysis["observations"]) == 28
    assert len(analysis["records"]) == 21
    assert len(analysis["comparisons"]) == 21
    assert all(r["eligible"] for r in analysis["records"])
    target = next(r for r in analysis["records"] if r["code"] == "00000001" and r["year"] == 2025)
    assert target["ratios"]["ar_turnover"] == D(1200000000) / D(395000000)
    assert target["ratios"]["gross_margin"] == D(".125")
    assert any("후보" in r["flag"] for r in analysis["cross_outliers"] if r["code"] == "00000001")
    assert all(r["rollforward_difference"] == 0 for r in analysis["records"])
    assert any("동반 상승" in f for r in analysis["records"] for f in r["credit_findings"])


def test_invalid_year_excluded_without_filling(inputs):
    snapshot, notes = deepcopy(inputs)
    snapshot["companies"][0]["years"]["2024"]["period"] = []
    result = analyze(snapshot, notes, "total_trade", D(".3"), D(5))
    rows = {r["year"]: r for r in result["records"] if r["code"] == "00000001"}
    assert not rows[2024]["eligible"]
    assert rows[2025]["ratios"]["ar_turnover"] is None
    assert rows[2025]["ratios"]["gross_margin"] is not None


def test_snapshot_replay_structure_and_identity(inputs, tmp_path):
    snapshot = inputs[0]
    path = tmp_path / "snapshot.json"
    save_snapshot(snapshot, path)
    assert load_snapshot(path) == snapshot
    bad = deepcopy(snapshot)
    bad["companies"][0]["years"]["2025"]["full"][0]["corp_code"] = "99999999"
    save_snapshot(bad, path)
    with pytest.raises(ValueError):
        load_snapshot(path)


def test_cli_demo_excel_png_and_replay(tmp_path):
    output = tmp_path / "demo"
    assert main(["--demo", "--output-dir", str(output)]) == 0
    path = output / "peer_financial_ratio_analysis_2025.xlsx"
    book = load_workbook(path, data_only=False)
    assert book.sheetnames == ["기업 기본정보", "원천 재무데이터", "재무비율 계산", "동종기업 비교", "이상치 분석", "시각화", "분석 방법 및 한계", "신용손실 검토", "계정매핑", "수치검증"]
    assert len(book["시각화"]._images) == 6
    assert book["재무비율 계산"]["E7"].value == pytest.approx(.125)
    assert "합성" in book["기업 기본정보"]["A2"].value
    assert len(book["동종기업 비교"].conditional_formatting) > 0
    assert all(cell.data_type != "f" for ws in book for row in ws for cell in row)
    book.close()
    images = list((output / "charts").glob("*.png"))
    assert len(images) == 6
    for image in images:
        with Image.open(image) as picture:
            assert picture.width >= 1500
            assert picture.height >= 800
    assert json.loads((output / "manifest.json").read_text(encoding="utf-8"))["source"] == "synthetic"
    replay = tmp_path / "replay"
    assert main(["--snapshot", str(output / "raw_snapshot.json"), "--credit-notes", str(output / "credit_notes.json"), "--output-dir", str(replay)]) == 0
    rerun = load_workbook(replay / path.name, data_only=True)
    assert rerun["재무비율 계산"]["E7"].value == pytest.approx(.125)
    rerun.close()


def test_report_formula_injection_is_plain_text(inputs, analysis, tmp_path):
    from excel_report import write_report
    snapshot = deepcopy(inputs[0])
    snapshot["companies"][0]["name"] = '=HYPERLINK("https://example.com")'
    path = write_report(tmp_path / "test.xlsx", snapshot, analysis, [], {})
    book = load_workbook(path)
    assert book["기업 기본정보"]["A5"].data_type == "s"
    assert book["기업 기본정보"]["A5"].value.startswith("=")
    book.close()


@pytest.mark.parametrize("args", [["--demo", "--turnover-change", "NaN"], ["--demo", "--start-year", "2024"], ["--demo", "--target", "실제회사"]])
def test_invalid_cli_input(args):
    assert main(args) == 2
