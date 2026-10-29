"""openpyxl로 근거 추적이 가능한 보고서와 PNG 삽입 시트를 작성한다."""
from decimal import Decimal
from pathlib import Path
import json
import os
import tempfile

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from ratio_calculator import METRICS

NAVY = "173F5F"
TEAL = "EAF3F5"
ORANGE = "FFF0E5"
NUM = '#,##0.00;[Red](#,##0.00);"-"'
MONEY = '#,##0;[Red](#,##0);"-"'
PCT = '0.0%;[Red](0.0%);"-"'


def write_report(path, snapshot, analysis, chart_paths, metadata):
    book = Workbook()
    book.remove(book.active)
    book.properties.title = "동종기업 재무비율 비교 및 이상치 검토"
    book.properties.description = "합성 테스트 데이터" if snapshot["synthetic"] else "OpenDART API 및 출처 있는 사용자 주석 자료"
    tag = "합성 TEST · 실제 공시 아님" if snapshot["synthetic"] else "OpenDART · 사용자 선택 비교기업"

    def table(name, headers, rows, formats=None):
        ws = book.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(2, len(headers)))
        ws.cell(1, 1, name).font = Font(name="맑은 고딕", size=18, bold=True, color=NAVY)
        ws.row_dimensions[1].height = 30
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(2, len(headers)))
        ws.cell(2, 1, f"{tag} | {snapshot['start_year']}–{snapshot['end_year']} | CFS / KRW 원 | 빈칸=계산 불가·미제공").font = Font(name="맑은 고딕", size=10, color="617482")
        ws.row_dimensions[2].height = 26
        for col, label in enumerate(headers, 1):
            cell = ws.cell(4, col, label)
            cell.font = Font(name="맑은 고딕", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(col)].width = 23
        ws.row_dimensions[4].height = 42
        for i, row in enumerate(rows, 5):
            ws.row_dimensions[i].height = 42
            for j, value in enumerate(row, 1):
                if isinstance(value, Decimal):
                    value = float(value)
                cell = ws.cell(i, j, value)
                # openpyxl의 = 문자열 자동 수식 판정을 차단하고 원문 문자열 유지.
                if isinstance(value, str):
                    cell.data_type = "s"
                cell.font = Font(name="맑은 고딕", size=10, color="263746")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = Border(bottom=Side(style="hair", color="D9E2E8"))
                if i % 2:
                    cell.fill = PatternFill("solid", fgColor="F3F7FA")
                if formats and j in formats:
                    cell.number_format = formats[j]
        ws.freeze_panes = "C5"
        ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{max(4, ws.max_row)}"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = "1:4"
        return ws

    companies = {c["code"]: c for c in snapshot["companies"]}
    info_rows = []
    for o in analysis["observations"]:
        c = companies[o["code"]]
        p = c["profile"]
        info_rows.append([c["name"], c["code"], c["stock_code"], c["role"], p.get("induty_code"), o["year"],
                          p.get("acc_mt"), o["period"]["balance_date"], str(o["period"]["income_dates"]),
                          "CFS", "KRW" if o["eligible"] else "검토 필요", "사용" if o["eligible"] else "제외",
                          "; ".join(o["reasons"]), "기초잔액 보조" if o["year"] < snapshot["start_year"] else "분석연도"])
    info = table("기업 기본정보", ["기업명", "고유번호", "종목코드", "역할", "현재 업종코드", "사업연도", "현재 결산월", "검증 결산일", "검증 손익기간", "요청 기준", "통화", "적격성", "제외 사유", "연도 역할"], info_rows)
    info.column_dimensions["M"].width = 85
    info.column_dimensions["I"].width = 48
    lookup = {(o["code"], o["year"]): o for o in analysis["observations"]}
    source_rows = []
    for r in analysis["records"]:
        previous = lookup[r["code"], r["year"] - 1]
        source_rows.append([r["name"], r["year"], r["values"]["revenue"], r["values"]["cost"], previous["credit"]["net"], r["credit"]["net"],
                            previous["values"]["inventory"], r["values"]["inventory"], ";".join(r["period"]["receipts"]),
                            ";".join(previous["period"]["receipts"]), r["credit"]["source_ref"],
                            "적격" if r["eligible"] else "비율 계산 제외; 원천만 보존"])
    sources = table("원천 재무데이터", ["기업명", "사업연도", "매출액(원)", "매출원가(원)", "기초 매출채권 순액", "기말 매출채권 순액", "기초 재고자산", "기말 재고자산", "당기 접수번호", "기초 접수번호", "매출채권 주석 근거", "적격성"], source_rows, {i: MONEY for i in range(3, 9)})
    sources.column_dimensions["K"].width = 60
    ratio_rows = []
    for r in analysis["records"]:
        ratio_rows.append([r["name"], r["role"], r["year"], *[r["ratios"][k] for k in METRICS],
                           "; ".join(f"{METRICS[k][0]}: {v}" for k, v in r["ratio_reasons"].items() if v)])
    ratios = table("재무비율 계산", ["기업명", "역할", "사업연도", *[f"{label} ({unit})" for label, unit in METRICS.values()], "계산 불가 사유"], ratio_rows,
                   {i + 4: PCT if METRICS[k][1] == "%" else NUM for i, k in enumerate(METRICS)})
    ratios.column_dimensions["L"].width = 100
    for i in range(5, ratios.max_row + 1):
        reason_text = ratios.cell(i, 12).value or ""
        ratios.row_dimensions[i].height = max(42, 15 * ((len(reason_text) + 49) // 50))

    comparison_rows = []
    for r in analysis["cross_outliers"]:
        comparison_rows.append([r["name"], r["role"], r["year"], METRICS[r["metric"]][0], r["value"], r["n"], r["mean"], r["median"],
                                r["mean_gap"], r["median_gap"], r["gap_unit"], r["rank"], r["flag"]])
    comp = table("동종기업 비교", ["기업명", "역할", "사업연도", "지표", "실제값", "유효 비교기업 n", "비교기업 평균", "비교기업 중앙값", "평균 대비 차이", "중앙값 대비 차이", "차이 단위", "내림차순 순위", "IQR 결과"], comparison_rows,
                 {i: NUM for i in (5, 7, 8, 9, 10)})
    for idx, record in enumerate(analysis["cross_outliers"], 5):
        if record["metric"] == "gross_margin":
            for col in (5, 7, 8):
                comp.cell(idx, col).number_format = PCT
    candidate_fill = PatternFill("solid", fgColor=ORANGE)
    comp.conditional_formatting.add(f"A5:M{max(5, comp.max_row)}", FormulaRule(formula=['ISNUMBER(SEARCH("후보",$M5))'], fill=candidate_fill))

    outlier_rows = []
    for r in analysis["cross_outliers"]:
        outlier_rows.append([r["name"], r["year"], METRICS[r["metric"]][0], "횡단면 IQR", r["value"], r["median"], r["n"], r["q1"], r["q3"], r["iqr"], r["lower"], r["upper"], None, None, "비율" if r["metric"] == "gross_margin" else "배", r["flag"], r["review"]])
    for r in analysis["time_outliers"]:
        outlier_rows.append([r["name"], r["year"], METRICS[r["metric"]][0], "시계열", r["value"], None, None, None, None, None, None, None, r["delta"], r["threshold"], r["unit"], r["flag"], r["review"]])
    outliers = table("이상치 분석", ["기업명", "사업연도", "지표", "분석유형", "실제값", "중앙값", "n", "Q1", "Q3", "IQR", "하한", "상한", "전년대비 변동", "변동 기준", "변동 단위", "판정", "추가 검토사항"], outlier_rows, {i: NUM for i in (5, 6, 8, 9, 10, 11, 12, 13, 14)})
    outliers.column_dimensions["Q"].width = 85
    outliers.conditional_formatting.add(f"A5:Q{max(5, outliers.max_row)}", FormulaRule(formula=['OR(ISNUMBER(SEARCH("후보",$P5)),ISNUMBER(SEARCH("검토",$P5)))'], fill=candidate_fill))
    for idx, r in enumerate(outlier_rows, 5):
        if r[2] == METRICS["gross_margin"][0]:
            for col in (5, 6, 8, 9, 10, 11, 12):
                outliers.cell(idx, col).number_format = PCT
        if r[14] == "비율" and r[3] == "시계열":
            outliers.cell(idx, 13).number_format = PCT
            outliers.cell(idx, 14).number_format = PCT

    visual = table("시각화", ["차트 / PNG 별도 저장"], [])
    visual.column_dimensions["A"].width = 30
    for idx, png in enumerate(chart_paths):
        picture = Image(str(png))
        picture.width, picture.height = 960, 470
        visual.add_image(picture, f"A{6 + idx * 27}")
    visual.auto_filter.ref = None
    visual.freeze_panes = None

    methods = [
        ("목적", "사용자 선택 동종기업 비교 및 추가 검토 신호. 통계적 이상치는 회계부정/오류 판단이 아님."),
        ("표본", "기준 평균·중앙값·IQR은 비교기업만 포함, 대상기업 제외. 비교기업 자신을 평가할 때는 기준집단에 자신도 포함."),
        ("순위", "대상기업+유효 비교기업 내 내림차순 경쟁순위(동률 1,1,3). 높은 값이 항상 좋은 것은 아님."),
        ("분석범위", "연결 CFS/KRW, 12월 결산 국내 상장 비금융. 기타채권/계약자산 통합 금액 미사용."),
        ("매출채권회전율", "매출액 / 평균 매출채권 순액. 기초·기말 순액과 범위를 주석으로 확인한 경우 계산."),
        ("매출총이익률", "(매출액 − 매출원가) / 매출액. 비율 자체는 %, 평균·중앙값과 차이는 %p."),
        ("재고자산회전율", "매출원가 / 평균 재고자산. 기초잔액용 직전 1개년 추가 수집."),
        ("회전기간", "365 / 회전율. 윤년에도 비교용 365일 관행 사용; 회전율 0은 기간 계산 불가."),
        ("IQR", "지표별 유효 비교기업 n≥5, IQR>0. Q1−1.5×IQR 미만 또는 Q3+1.5×IQR 초과. 경계값은 후보 아님."),
        ("사분위수", "선형 보간 h=(n−1)p, PERCENTILE.INC와 같은 정의."),
        ("시계열 기준", f"회전율 상대변화 ±{analysis['turnover_threshold'] * 100}% 초과; 이익률 ±{analysis['margin_pp']}%p 초과."),
        ("주석 범위", analysis["scope"] + ": 모든 기업에 동일하게 적용. current_trade와 total_trade를 혼용하지 않음."),
        ("손실충당금", "기말 잔액(총액 대비 설정률)과 당기 ECL 비용/환입을 구분. 순액에서 충당금을 재차 차감하지 않음."),
        ("총액 회전율", "주석의 총액 기초·기말이 있을 때 별도 표시. 제각/양도도 회전율 개선처럼 보일 수 있어 회수 증거 아님."),
        ("충당금 롤포워드", "기초 충당금 + ECL 비용(환입은 음수) − 제각 + 기타변동 − 기말 충당금. 필수항목 없으면 대사 유보. 회수액은 자동 가산하지 않음."),
        ("ECL 한계", "매출채권만의 출처 있는 공시 금액을 사용. 노출별 연체·전망·손실률 없이 ECL 적정성을 추정/판단하지 않음."),
        ("시점", "각 연도 API 당기금액을 사용. 후속 보고서 재작성 비교금액으로 자동 소급하지 않음. 기업개황은 조회시점 정보."),
        ("비교 한계", "신용매출 대신 전체 매출 사용, 계절성·매출채권 양도·사업결합·재고평가손실·제품구성·회계정책 차이 미조정."),
        ("비교 한계 2", "비교기업 집단은 업종 전체 평균이 아님. 기준집단에 포함된 이상치를 자동 제거하지 않음. 연도별 결측으로 표본이 다를 수 있음."),
        ("통화/기간", "기업개황 acc_mt만으로 과거 결산일을 추정하지 않음. fnlttSinglAcnt.json의 동일 공시 CFS 날짜로 1년 기간 검증."),
        ("데이터 출처", "OpenDART corpCode.xml, company.json, fnlttSinglAcntAll.json, fnlttSinglAcnt.json + 출처 있는 사용자 주석 입력"),
        ("재계산", "모든 계산은 Python 수행. Excel은 스냅샷이며 셀 수정 후 자동 재계산되지 않음."),
        ("정밀도", "Python Decimal 계산 후 Excel 숫자로 변환. 15자리 초과 원화 원문은 raw_snapshot.json 확인."),
        ("매핑", "계정매핑 시트에 선택값 및 모든 후보 기록. 같은 계정 중복 합산 없음. IS 우선, CIS 보조."),
        ("재현", "같은 스냅샷·주석·옵션으로 재실행. manifest.json의 입력 SHA-256과 라이브러리 버전 확인."),
        *[(str(k), str(v)) for k, v in metadata.items()],
    ]
    method = table("분석 방법 및 한계", ["항목", "내용"], methods)
    method.column_dimensions["B"].width = 130
    for i in range(5, method.max_row + 1):
        method.row_dimensions[i].height = 48

    credit_rows = []
    for r in analysis["records"]:
        credit = r["credit"]
        previous = lookup[r["code"], r["year"] - 1]["credit"]
        credit_rows.append([r["name"], r["year"], analysis["scope"], credit["gross"], credit["allowance"], credit["net"], credit["reconciliation"],
                            previous.get("allowance"), credit["ecl_expense"], credit["writeoffs"], credit["recoveries"], credit["other_movements"], r["rollforward_difference"],
                            r["ratios"]["allowance_ratio"], r["ratios"]["gross_ar_turnover"], credit["status"], credit["reason"],
                            "; ".join(r["credit_findings"]), credit["source_ref"], credit["source_url"]])
    cr = table("신용손실 검토", ["기업명", "사업연도", "채권 범위", "매출채권 총액", "기말 손실충당금", "매출채권 순액", "총액-충당금-순액", "기초 충당금", "ECL 비용/환입", "제각", "회수액(별도)", "기타 충당금 변동", "충당금 대사 차이", "충당금 설정률", "총액 회전율(배)", "주석 확인 상태", "확인 사유", "검토 신호", "주석 위치", "출처 URL"], credit_rows,
               {**{i: MONEY for i in range(4, 14)}, 14: PCT, 15: NUM})
    for col in ("Q", "R", "S", "T"):
        cr.column_dimensions[col].width = 65
    for i in range(5, cr.max_row + 1):
        cr.row_dimensions[i].height = 60

    mapping_rows, check_rows = [], []
    for o in analysis["observations"]:
        for item, m in o["mapping"].items():
            mapping_rows.append([o["name"], o["year"], item, m.amount, m.account_id, m.account_nm, m.statement, m.currency, m.receipt,
                                 m.method, m.status, m.reason, json.dumps(m.candidates, ensure_ascii=False),
                                 f"https://dart.fss.or.kr/dsaf001/main.do?rcpNo={m.receipt}" if m.receipt else ""])
        for check in o["checks"]:
            check_rows.append([o["name"], o["year"], check["check"], check["difference"], check["status"], "허용오차 1원; PASS는 감사의견 아님"])
    mapping = table("계정매핑", ["기업명", "사업연도", "항목", "선택 금액(원)", "계정 ID", "계정명", "재무제표", "통화", "접수번호", "선택방법", "상태", "사유", "모든 후보(JSON)", "공시 링크"], mapping_rows, {4: MONEY})
    mapping.column_dimensions["M"].width = 110
    mapping.column_dimensions["E"].width = 48
    checks = table("수치검증", ["기업명", "사업연도", "검증", "차이(원)", "상태", "기준"], check_rows, {4: MONEY})
    checks.conditional_formatting.add(f"A5:F{max(5, checks.max_row)}", FormulaRule(formula=['$E5="FAIL"'], fill=candidate_fill))

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(suffix=".xlsx", dir=path.parent)
    os.close(descriptor)
    try:
        book.save(temporary)
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
        book.close()
    return path
